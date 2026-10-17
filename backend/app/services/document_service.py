import io
from pathlib import Path
from typing import BinaryIO, Union

import docx
import pypdf

from app.core.config import settings


class DocumentExtractionError(Exception):
    """Custom exception raised when document extraction fails."""
    pass


class DocumentService:
    """
    Handles file ingestion, format validation, and text extraction for TXT, PDF, and DOCX documents.
    """

    SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}

    @classmethod
    def validate_file(cls, filename: str, file_size: int) -> str:
        """
        Validates file extension and size constraints.
        Returns the normalized lowercase file extension.
        """
        if not filename:
            raise DocumentExtractionError("No filename provided.")

        ext = Path(filename).suffix.lower()
        if ext not in cls.SUPPORTED_EXTENSIONS:
            allowed = ", ".join(cls.SUPPORTED_EXTENSIONS)
            raise DocumentExtractionError(
                f"Unsupported file format '{ext}'. Supported formats are: {allowed}."
            )

        if file_size > settings.MAX_UPLOAD_SIZE_BYTES:
            max_mb = settings.MAX_UPLOAD_SIZE_BYTES // (1024 * 1024)
            raise DocumentExtractionError(
                f"File size exceeds maximum allowed upload limit of {max_mb} MB."
            )

        return ext

    @classmethod
    def extract_text_from_file(
        cls,
        filename: str,
        file_bytes: bytes,
    ) -> str:
        """
        Extracts clean plain text from uploaded file bytes.
        """
        if not file_bytes:
            raise DocumentExtractionError("Uploaded file is empty.")

        ext = cls.validate_file(filename, len(file_bytes))

        if ext == ".txt":
            return cls._extract_txt(file_bytes)
        elif ext == ".pdf":
            return cls._extract_pdf(file_bytes)
        elif ext == ".docx":
            return cls._extract_docx(file_bytes)
        else:
            raise DocumentExtractionError(f"Unsupported document format: {ext}")

    @classmethod
    def _extract_txt(cls, file_bytes: bytes) -> str:
        """Decodes raw text bytes using UTF-8 with fallback to latin-1."""
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                return file_bytes.decode("latin-1")
            except Exception as e:
                raise DocumentExtractionError(f"Failed to decode text file: {e}")

    @classmethod
    def _extract_pdf(cls, file_bytes: bytes) -> str:
        """Extracts text from PDF pages using pypdf."""
        try:
            stream = io.BytesIO(file_bytes)
            reader = pypdf.PdfReader(stream)
            extracted_pages = []

            for page_num, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    extracted_pages.append(text)

            full_text = "\n\n".join(extracted_pages).strip()
            if not full_text:
                raise DocumentExtractionError(
                    "The uploaded PDF contains no extractable text. Scanned or image-based "
                    "PDFs require OCR, which is outside the initial scope."
                )

            return full_text
        except DocumentExtractionError:
            raise
        except Exception as e:
            raise DocumentExtractionError(f"Failed to parse PDF document: {str(e)}")

    @classmethod
    def _extract_docx(cls, file_bytes: bytes) -> str:
        """Extracts text from DOCX paragraphs and tables using python-docx."""
        try:
            stream = io.BytesIO(file_bytes)
            doc = docx.Document(stream)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

            # Also inspect table cells
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text.strip())

            full_text = "\n\n".join(paragraphs).strip()
            if not full_text:
                raise DocumentExtractionError("The uploaded DOCX document contains no readable text.")

            return full_text
        except DocumentExtractionError:
            raise
        except Exception as e:
            raise DocumentExtractionError(f"Failed to parse DOCX document: {str(e)}")


document_service = DocumentService()
